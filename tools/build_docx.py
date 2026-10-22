#!/usr/bin/env python3
"""Convert the client proposal (markdown) into a Word document in INT's proposal style.

Usage:
    python build_docx.py Proposal_Client_Title_V1.0.md -o Proposal_Client_Title_V1.0.docx [--version V1.0]

Input conventions (the same markdown the master writes and the scripted checks read):
    * everything before the first `---` line (or, without one, before the first `##` heading) is the cover page: `# Title`, a `**subtitle**` line, plain lines, `Presented by INT.`, `*tagline*`
    * `## 1. Section` becomes Heading 1, `###` Heading 2, `####` Heading 3
    * `- item` / `* item` bullets (indent two spaces per level), `1. item` numbered items, `| a | b |` tables, `> quote`, ``` code fences
    * **bold**, *italic* and `code` inside text; the first `---` is a page break (end of the cover), later ones are ignored
Layout: A4, Calibri 10.5, navy headings (1F3864), navy table headers with white text, running header with the title and version,
footer "Presented by INT | Confidential | Page N".
"""
import argparse
import os
import re
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit("python-docx is required: pip install python-docx")

NAVY = RGBColor(0x1F, 0x38, 0x64)
TITLE_NAVY = RGBColor(0x17, 0x36, 0x5D)
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*(?!\s)[^*]+?(?<!\s)\*(?!\*))")


def add_runs(par, text, size=None, bold=None, italic=None, color=None):
    """Add text to a paragraph, honouring **bold**, *italic* and `code`."""
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, mono = bold, italic, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, mono = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        run = par.add_run(part)
        if size:
            run.font.size = Pt(size)
        if b is not None:
            run.bold = b
        if i is not None:
            run.italic = i
        if color is not None:
            run.font.color.rgb = color
        if mono:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_page_field(par, size=9):
    run = par.add_run()
    run.font.size = Pt(size)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def setup(doc, header_text):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = "Calibri", Pt(size), True, NAVY
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(14 if name == "Heading 1" else 10), Pt(6)
        st.paragraph_format.keep_with_next = True
    hp = sec.header.paragraphs[0]
    hp.text = ""
    add_runs(hp, header_text, size=9, color=RGBColor(0x59, 0x59, 0x59))
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    add_runs(fp, "Presented by INT  |  Confidential  |  Page ", size=9, color=RGBColor(0x59, 0x59, 0x59))
    add_page_field(fp)


def split_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip().replace("\\|", "|") for c in cells]


def is_table_sep(line):
    return bool(re.fullmatch(r"\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*", line)) and "-" in line


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            text = row[ci] if ci < len(row) else ""
            if ri == 0:
                shade(cell, "1F3864")
                add_runs(par, text, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                add_runs(par, text, size=9.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


def build(md_text, version="", out_stats=None):
    lines = md_text.replace("\r\n", "\n").split("\n")
    title = next((l[2:].strip() for l in lines if l.startswith("# ")), "Proposal")
    if not version:
        m = re.search(r"^\s*Version:\s*(\S+)", md_text, re.M)
        version = m.group(1) if m else ""
    header = f"{title}  |  Proposal {version}".rstrip()
    doc = Document()
    setup(doc, header)

    i, n = 0, len(lines)
    in_cover, seen_rule, cover_started = True, False, False
    stats = {"headings": 0, "tables": 0, "bullets": 0, "paragraphs": 0}
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            if in_cover and not seen_rule:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                in_cover, seen_rule = False, True
            i += 1
            continue
        if in_cover and re.match(r"#{2,6}\s", stripped):            # no rule after the cover: the first section heading ends it
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            in_cover, seen_rule = False, True
        if in_cover:
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if stripped.startswith("# "):
                par.paragraph_format.space_before = Pt(110)
                par.paragraph_format.space_after = Pt(14)
                add_runs(par, stripped[2:].strip(), size=28, bold=True, color=TITLE_NAVY)
            elif stripped.strip("*_ ").lower().startswith("presented by"):
                par.paragraph_format.space_before = Pt(30)
                add_runs(par, stripped.strip("*_ "), size=13, bold=True)
            elif stripped.strip("*_ ").lower().startswith("we make it possible"):
                add_runs(par, stripped.strip("*_ "), size=12, italic=True)
            elif stripped.startswith("**") and stripped.endswith("**"):
                par.paragraph_format.space_after = Pt(18)
                add_runs(par, stripped.strip("*"), size=16, bold=True)
            elif stripped.startswith("*") and stripped.endswith("*"):
                add_runs(par, stripped.strip("*"), size=12, italic=True)
            else:
                add_runs(par, stripped, size=11)
            cover_started = True
            i += 1
            continue
        m = re.match(r"(#{1,6})\s+(.*)", stripped)
        if m:
            marks = len(m.group(1))                      # "#" and "##" -> Heading 1, "###" -> Heading 2, "####" -> Heading 3
            doc.add_heading(m.group(2).strip(), level=1 if marks <= 2 else min(marks - 1, 3))
            stats["headings"] += 1
            i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            par = doc.add_paragraph()
            run = par.add_run("\n".join(block))
            run.font.name, run.font.size = "Consolas", Pt(9)
            continue
        if stripped.startswith("|") and i + 1 < n and is_table_sep(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            stats["tables"] += 1
            continue
        mb = re.match(r"(\s*)([-*+])\s+(.*)", line)
        if mb:
            level = min(len(mb.group(1).replace("\t", "  ")) // 2, 2)
            par = doc.add_paragraph(style="List Bullet" if level == 0 else f"List Bullet {level + 1}")
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, mb.group(3))
            stats["bullets"] += 1
            i += 1
            continue
        mn = re.match(r"(\s*)(\d+)[.)]\s+(.*)", line)
        if mn:
            # literal numbers with a hanging indent: Word's automatic numbering would continue across separate lists
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.6)
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, f"{mn.group(2)}. {mn.group(3)}")
            stats["bullets"] += 1
            i += 1
            continue
        if stripped.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.8)
            add_runs(par, stripped.lstrip("> ").strip(), italic=True)
            i += 1
            continue
        par = doc.add_paragraph()
        add_runs(par, stripped)
        stats["paragraphs"] += 1
        i += 1
    if out_stats is not None:
        out_stats.update(stats, title=title, version=version, cover=cover_started)
    return doc


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--version", default="")
    a = ap.parse_args()
    with open(a.markdown, encoding="utf-8") as f:
        text = f.read()
    stats = {}
    doc = build(text, a.version, stats)
    doc.save(a.out)
    print(f"wrote {a.out}: {stats['headings']} headings, {stats['tables']} tables, {stats['bullets']} list items, {stats['paragraphs']} paragraphs")


if __name__ == "__main__":
    main()
