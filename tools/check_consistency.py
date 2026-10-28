#!/usr/bin/env python3
"""Scripted artifact-consistency checks (X-checks) for one run's output folder.

Expected files in --dir (first match wins):
    estimation_revision*.json, else estimation_pass2_retry*.json, else estimation_pass2.json (newest COMPLETE estimate; or --estimate NAME)
    estimation_derived.json                             (calculator output for that estimate)
    Estimation_*.xlsx    Proposal_*.md    Presales_Review_Notes_*.md
    case.json (optional)  -> X10 hash verification

Usage:
    python check_consistency.py --dir output/Client_20261012 [--forbid "NameA,NameB"] [--expose-hours]
Exit code 1 if any check fails. Writes consistency_report.json into --dir.
"""
import argparse
import glob
import hashlib
import json
import math
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import estimate_calc as ec  # noqa: E402

BLOCKING = {"S1", "S2", "S3", "S4", "S5", "T1", "T2", "T3", "T4", "P1", "P2", "A1", "A2", "A3", "A4", "O1", "R1"}
STOP = {"with", "that", "this", "from", "will", "have", "been", "were", "their", "which", "into", "than", "when", "each",
        "other", "such", "also", "must", "shall", "should", "would", "could", "about", "after", "before", "within"}
LEAK_TERMS = ["presales review notes", "sub-agent", "subagent", "estimation agent", "calculator", "pert", "most likely",
              "optimistic", "pessimistic", "contingency %", "confidence:", "effort exposure"]


def words(s):
    return {w for w in re.findall(r"[a-z0-9]+", s.lower()) if len(w) > 3 and w not in STOP}


def covered(needle, hay, thr=0.5):
    nw = words(needle)
    return True if not nw else len(nw & words(hay)) / len(nw) >= thr


def first(dirpath, *patterns):
    for p in patterns:
        m = sorted(glob.glob(os.path.join(dirpath, p)))
        if m:
            return m[0]
    return None


def usable_estimate(path):
    """A complete estimate package: valid JSON object with modules that hold tasks (not a patch or a cut-off file)."""
    try:
        e = ec.load(path)
    except (ValueError, OSError):
        return False
    return isinstance(e, dict) and any(isinstance(m, dict) and m.get("tasks") for m in e.get("modules") or [])


def pick_estimate(dirpath):
    """Latest complete estimate: revision, then the newest pass-2 retry, then pass 2. Skips unusable files."""
    cands = sorted(glob.glob(os.path.join(dirpath, "estimation_revision*.json")), reverse=True)         + sorted(glob.glob(os.path.join(dirpath, "estimation_pass2_retry*.json")), reverse=True)         + sorted(glob.glob(os.path.join(dirpath, "estimation_pass2.json")))
    return next((c for c in cands if usable_estimate(c)), None)


def sections(md):
    """Return {section_number: text} using numbered markdown headings like '## 8. Functional ...'."""
    heads = [(m.start(), int(m.group(1))) for m in re.finditer(r"^#{1,4}\s*(\d{1,2})[\.\)]?\s+\S.*$", md, re.M)]
    out = {}
    for i, (pos, num) in enumerate(heads):
        end = heads[i + 1][0] if i + 1 < len(heads) else len(md)
        out.setdefault(num, md[pos:end])
    return out


def subsection(text, name):
    """Text of the `###`/`####` subsection whose heading contains `name` (case-insensitive), up to the next heading of the same or higher level."""
    m = re.search(r"^(#{3,4})[ \t]*[^\n]*" + re.escape(name) + r"[^\n]*$", text, re.M | re.I)
    if not m:
        return ""
    lvl = len(m.group(1))
    nxt = re.search(r"^#{1," + str(lvl) + r"}[ \t]", text[m.end():], re.M)
    return text[m.start():m.end() + (nxt.start() if nxt else len(text) - m.end())]


def md_tables(text):
    """Every markdown table in `text` as a list of rows (a list of cell strings); the header row first, the separator row dropped."""
    out, cur = [], []
    for line in text.split("\n") + [""]:
        if line.strip().startswith("|"):
            cells = [c.strip().replace("\\|", "|") for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                cur.append(cells)
        elif cur:
            out.append(cur)
            cur = []
    return out


def plain(s):
    return re.sub(r"[*_`]", "", s).strip().lower()


def paragraphs(sec_text):
    """Prose paragraphs of a section (headings, bullets, tables and pictures are not counted)."""
    body = "\n".join(sec_text.split("\n")[1:])
    out = []
    for block in re.split(r"\n\s*\n", body):
        b = block.strip()
        if b and not re.match(r"(#|\||[-*+]\s|\d+[.)]\s|!\[|>)", b):
            out.append(b)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dir", required=True)
    ap.add_argument("--forbid", default="", help="comma-separated names that must not appear in the client proposal")
    ap.add_argument("--expose-hours", action="store_true", help="allow person-hours in the proposal")
    ap.add_argument("--estimate", default="", help="file name of the final estimate in --dir (default: newest complete revision / retry / pass 2)")
    ap.add_argument("--derived", default="", help="file name of the calculator output in --dir (default: estimation_derived.json)")
    a = ap.parse_args()
    d = a.dir

    est_p = os.path.join(d, a.estimate) if a.estimate else pick_estimate(d)
    der_p = os.path.join(d, a.derived) if a.derived else first(d, "estimation_derived.json")
    est_p = est_p if est_p and os.path.exists(est_p) else None
    der_p = der_p if der_p and os.path.exists(der_p) else None
    xl_p = first(d, "Estimation_*.xlsx")
    prop_p = first(d, "Proposal_*.md")
    notes_p = first(d, "Presales_Review_Notes_*.md")
    case_p = first(d, "case.json")
    res = []

    def chk(cid, ok, detail, level="fail"):
        res.append({"id": cid, "pass": bool(ok), "detail": detail, "level": level if not ok else "pass"})

    missing = [n for n, p in (("estimate", est_p), ("derived", der_p), ("proposal", prop_p)) if not p]
    if missing:
        chk("X0", False, f"required files not found in {d}: {missing}")
        est = der = None
    else:
        est, der = None, None
        try:
            est, der = ec.load(est_p), json.load(open(der_p, encoding="utf-8"))
        except ValueError as e:
            chk("X0", False, f"cannot read {os.path.basename(est_p)} or {os.path.basename(der_p)} as JSON: {e}")
        if est is not None and not usable_estimate(est_p):
            chk("X0", False, f"{os.path.basename(est_p)} is not a complete estimate (no modules with tasks): pass the right file as `estimate`")
            est = None
    prop = open(prop_p, encoding="utf-8").read() if prop_p else ""
    notes = open(notes_p, encoding="utf-8").read() if notes_p else ""
    secs = sections(prop)

    if est and der:
        # X1 workbook vs derived (independent recomputation from cells)
        if xl_p:
            try:
                import build_workbook as bw
                from openpyxl import load_workbook
                ws = load_workbook(xl_p)[bw.EB]
                n = base = 0.0
                for row in ws.iter_rows(min_row=bw.FIRST, max_col=7):
                    if isinstance(row[0].value, (int, float)):
                        o, p, m = (float(row[i].value) for i in (4, 5, 6))
                        base, n = base + (o + 4 * m + p) / 6, n + 1
                pct = sum(float(ws.cell(r, 11).value or 0) for r in (bw.R_QA, bw.R_PM, bw.R_UAT, bw.R_DOC, bw.R_CONT))
                tot = base * (1 + pct / 100)
                ok = int(n) == der["inputs"]["task_count"] and abs(tot - der["total"]["pert"]) <= 0.15
                chk("X1", ok, f"workbook tasks {int(n)} vs {der['inputs']['task_count']}; total {tot:.1f} vs {der['total']['pert']}")
                banner = str(ws["A1"].value or "")
                failed = [c["id"] for c in der["checks"] if not c["pass"] and c["id"] in BLOCKING]
                chk("X12", (not failed) or banner.startswith("DRAFT"),
                    f"blocking calculator failures {failed}; workbook banner: {banner[:50]!r}")
            except ImportError:
                chk("X1", False, "openpyxl not installed; cannot read workbook", "warn")
        else:
            chk("X1", False, "no Estimation_*.xlsx found")

        # X2 modules in Section 8
        s8 = secs.get(8, "")
        gone = [m["module"] for m in est["modules"] if not covered(m["module"], s8 or prop, 0.6)]
        chk("X2", not gone, f"estimated modules not found in Section 8: {gone}" if gone else "all estimated modules appear in scope")

        # X3 weeks in Section 18
        s18 = secs.get(18, "")
        nums = [float(x) for m in re.finditer(r"(\d+(?:\.\d+)?)(?:\s*(?:–|-|to)\s*(\d+(?:\.\d+)?))?\s*weeks?", s18, re.I)
                for x in m.groups() if x]
        lo, hi = math.floor(der["timeline"]["weeks_min"]) - 1, math.ceil(der["timeline"]["weeks_max"]) + 1
        if not nums:
            chk("X3", False, "no week figures found in Section 18 (placeholder?)", "warn")
        else:
            out = [x for x in nums if not lo <= x <= hi]
            chk("X3", not out, f"Section 18 weeks {nums}; allowed {lo}-{hi}" + (f"; outside: {out}" if out else ""))

        # X4 tech stack in Section 10
        s10 = secs.get(10, "")
        nt = [t["proposed"] for t in est.get("tech_stack", []) if not covered(t["proposed"], s10 or prop, 0.5)]
        chk("X4", not nt, f"stack items not in Section 10: {nt}" if nt else "stack matches Section 10")

        # X5 assumptions and dependencies
        s17, s16 = secs.get(17, ""), secs.get(16, "")
        ua = [x["text"] for x in est.get("assumptions", []) if x.get("source") == "unanswered_question"]
        ua += [g["assumption"] for g in est.get("gaps_assumed", [])]
        miss_a = [t[:60] for t in ua if not covered(t, s17 or prop)]
        miss_d = [x["dependency"][:60] for x in est.get("client_dependencies", []) if not covered(x["dependency"], s16 or prop)]
        chk("X5", not miss_a and not miss_d, f"missing assumptions {miss_a}; missing dependencies {miss_d}")

    # X6 money
    money = [m.group(0) for m in re.finditer(r"(?:₹|\bINR\b|\bRs\.?|\bUSD\b|\$)\s?[\d][\d,\.]*", prop)]
    chk("X6", not money, f"amounts in proposal: {money[:5]}" if money else "no amounts other than placeholders")
    chk("X6b", "XXXXX" in prop, "price placeholder XXXXX present in Section 19", "warn")

    # X7 leaks
    low = prop.lower()
    leaks = [t for t in LEAK_TERMS if t in low] + re.findall(r"\bR-\d+\b", prop)
    chk("X7", not leaks, f"internal terms in client proposal: {leaks}" if leaks else "no internal content leaked")

    # X8 milestones
    ms_text = secs.get(19, prop)
    pcts = [float(x) for x in re.findall(r"\|\s*\**\s*(\d+(?:\.\d+)?)\s*%\s*\**\s*\|", ms_text)]
    body = pcts[:-1] if pcts and abs(pcts[-1] - 100) < 0.01 and len(pcts) > 1 else pcts
    okm = bool(body) and abs(sum(body) - 100) < 0.01 and body[0] == 20 and body[-1] == 15
    chk("X8", okm, f"milestone percentages {pcts}")

    # X9 sections
    nums = sorted(secs)
    miss = [i for i in range(1, 23) if i not in secs]
    chk("X9", not miss, f"missing section numbers: {miss}" if miss else "all 22 sections present")

    # X16-X20 content elements every proposal must carry (see master/proposal_template.md section 8)
    n16 = len(paragraphs(secs.get(1, "")))
    chk("X16", 3 <= n16 <= 4, f"Executive Summary has {n16} paragraph(s); it must have 3 to 4")

    s2 = secs.get(2, "")
    gaps = []
    for label in ("Platform Goals", "Business Objectives"):
        sub = subsection(s2, label)
        k = len(re.findall(r"^\s*[-*+]\s+\S", sub, re.M))
        if not sub:
            gaps.append(f"no '{label}' subsection in Section 2")
        elif not 3 <= k <= 7:
            gaps.append(f"'{label}' has {k} item(s); use 3 to 7")
    chk("X17", not [g for g in gaps if g.startswith("no ")], "; ".join(gaps) if gaps else "Platform Goals and Business Objectives present")
    if any(not g.startswith("no ") for g in gaps):
        chk("X17b", False, "; ".join(g for g in gaps if not g.startswith("no ")), "warn")

    sub4 = subsection(secs.get(4, ""), "Current State Assessment")
    map_tab = next((t for t in md_tables(sub4) if any("proposed solution" in plain(c) for c in t[0])), None)
    if not map_tab or len(map_tab) < 2:
        chk("X18", False, "Section 4 needs a 'Current State Assessment' subsection with a table (Current-state finding | Source | Proposed solution | Scope reference)")
    else:
        head = [plain(c) for c in map_tab[0]]
        ci = next((i for i, c in enumerate(head) if "scope" in c or "module" in c), len(head) - 1)
        si = next((i for i, c in enumerate(head) if c == "source"), None)
        s8m = secs.get(8, "")
        unmapped = [r[0][:50] for r in map_tab[1:] if len(r) > ci and not covered(r[ci], s8m, 0.6)]
        badsrc = [r[0][:40] for r in map_tab[1:] if si is not None and len(r) > si and plain(r[si]) not in ("client-stated", "observed", "assumption")]
        chk("X18", True, f"{len(map_tab) - 1} current-state row(s) mapped to solutions")
        if unmapped:
            chk("X18b", False, f"rows whose scope reference is not found in Section 8: {unmapped}", "warn")
        if badsrc:
            chk("X18c", False, f"rows without a Source of Client-stated, Observed or Assumption: {badsrc}", "warn")

    roles = []
    for t in md_tables(subsection(secs.get(8, ""), "Roles")):
        if re.search(r"role|user", plain(t[0][0])):
            roles += [plain(r[0]) for r in t[1:] if r and plain(r[0])]
    rbac = subsection(secs.get(13, ""), "Role-Based Access Control") or subsection(secs.get(13, ""), "RBAC")
    mat = next((t for t in md_tables(rbac) if re.search(r"role", plain(t[0][0]))), None)
    if not roles:
        chk("X19", False, "Section 8 needs a roles subsection ('User Types, Roles and Responsibilities') with a table whose first column is the role")
    elif not mat:
        chk("X19", False, "Section 13 needs a 'Role-Based Access Control' subsection with a role x module matrix (first column: role)")
    else:
        in_mat = [plain(r[0]) for r in mat[1:] if r]
        lost = [r for r in roles if not any(r in m or m in r for m in in_mat)]
        chk("X19", not lost, f"roles missing from the RBAC matrix: {lost}" if lost else f"{len(roles)} role(s) all appear in the RBAC matrix")
        if "proposed" not in rbac.lower():
            chk("X19b", False, "the RBAC matrix is not labelled Proposed", "warn")

    pics = list(re.finditer(r"^!\[(.*?)\]\(([^()\s]+)\)\s*$", prop, re.M))
    if not pics:
        chk("X20", False, "no user-journey diagram in the proposal", "warn")
    else:
        bad = []
        for m in pics:
            cap, fn = m.group(1).strip(), m.group(2)
            after = [l for l in prop[m.end():].split("\n") if l.strip()][:4]
            if not os.path.exists(os.path.join(d, os.path.basename(fn))):
                bad.append(f"{fn}: file not found")
            elif not cap:
                bad.append(f"{fn}: no caption")
            elif not any(l.strip().startswith("|") and "step" in l.lower() for l in after):
                bad.append(f"{fn}: no step table under it")
        chk("X20", not bad, "; ".join(bad) if bad else f"{len(pics)} diagram(s) with caption and step table", "fail")

    # X10 case-file hashes
    if case_p:
        case = json.load(open(case_p, encoding="utf-8"))
        bad = []
        for name, f in case.get("files", {}).items():
            if not os.path.exists(f["path"]):
                bad.append(f"{name} missing")
            else:
                h = hashlib.sha256(open(f["path"], "rb").read()).hexdigest()
                if h != f["sha256"]:
                    bad.append(f"{name} changed")
        chk("X10", not bad, "; ".join(bad) if bad else "registered files unchanged")
    else:
        chk("X10", False, "no case.json; hashes not verified", "warn")

    # X11 review notes list uncovered items and inferred tasks
    if der:
        need = list(der.get("uncovered_requirement_items", [])) + list(der.get("inferred_tasks", []))
        absent = [i for i in need if i not in notes]
        chk("X11", bool(notes_p) and not absent, f"review notes missing: {absent}" if notes_p else "no review notes file")

    # X13 forbidden names
    names = [n.strip() for n in a.forbid.split(",") if n.strip()]
    est_text = json.dumps(est) if est else ""
    hit = [n for n in names if n.lower() in low or n.lower() in est_text.lower()]
    chk("X13", not hit, f"forbidden names present: {hit}" if hit else "no forbidden client names", "fail")

    # X15 hours in proposal
    if not a.expose_hours:
        hrs = re.findall(r"\d[\d,]*(?:\.\d+)?\s*(?:person[- ]?hours|man[- ]?hours|person[- ]?days|man[- ]?days|hrs|hours)", prop, re.I)
        chk("X15", not hrs, f"person-hours in client proposal: {hrs[:5]}" if hrs else "no person-hours in proposal")

    fails = [r for r in res if not r["pass"] and r["level"] == "fail"]
    warns = [r for r in res if not r["pass"] and r["level"] == "warn"]
    for r in res:
        tag = "PASS" if r["pass"] else ("WARN" if r["level"] == "warn" else "FAIL")
        print(f"{tag} {r['id']}: {r['detail']}")
    print(f"{len(res) - len(fails) - len(warns)} pass, {len(warns)} warn, {len(fails)} fail")
    with open(os.path.join(d, "consistency_report.json"), "w", encoding="utf-8") as f:
        json.dump(res, f, indent=2, ensure_ascii=False)
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
