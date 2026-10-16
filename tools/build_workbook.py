#!/usr/bin/env python3
"""Build the estimation workbook (.xlsx) from the estimator's JSON and the calculator's output.

Layout mirrors INT's UWS estimate: Estimation Breakdown (+ summary block), Assumptions,
Estimation Basis, plus Effort by Type, Team & Timeline, Tech Stack, Risks, Cost Drivers,
Maintenance, Clarifications Log, Requirement Coverage.
Inputs (task text, O/P/M, percentages) come from the estimate JSON; derived numbers that are not
live formulas (team hours, phases, weeks, P80) come from the calculator output. Row 1 of every
sheet is a banner showing the estimate state or a DRAFT stamp.

Usage:
    python build_workbook.py estimation_pass2.json estimation_derived.json -o Estimation_Client_V1.0.xlsx \
        [--client NAME] [--version 1.0] [--state draft|delivery_reviewed|approved] \
        [--draft "S4,T1"] [--clarifications clarifications.json] [--sources "a.pdf,b.pdf"] [--check]
--check re-reads the file and recomputes the totals independently from its cells.
"""
import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import estimate_calc as ec  # noqa: E402

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

EB = "Estimation Breakdown"
HDR, FIRST = 2, 3
# summary block rows on the Estimation Breakdown sheet
R_DEV, R_QA, R_PM, R_UAT, R_DOC, R_CONT, R_TOT = 3, 4, 5, 6, 7, 8, 9
R_O, R_M, R_P, R_P80, R_STATE, R_HPW = 11, 12, 13, 14, 15, 16
LABELS = {"qa": "QA + Application Security Validation", "pm": "PM Effort", "uat": "UAT",
          "docs": "Documentation & Knowledge Transfer", "cont": "Contingency"}
ROW_OF = {"qa": R_QA, "pm": R_PM, "uat": R_UAT, "docs": R_DOC, "cont": R_CONT}
BOLD = Font(bold=True)
HFILL = PatternFill("solid", fgColor="DDE6F2")
BFILL = PatternFill("solid", fgColor="F2F2F2")
WRAP = Alignment(wrap_text=True, vertical="top")


def banner(ws, text, ncols=8):
    ws.cell(1, 1, text).font = Font(bold=True, color="9C0006" if text.startswith("DRAFT") else "1F4E78")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)


def header(ws, row, names):
    for i, n in enumerate(names, 1):
        c = ws.cell(row, i, n)
        c.font, c.fill, c.alignment = BOLD, HFILL, WRAP


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def simple_sheet(wb, title, banner_text, cols, rows, col_widths):
    ws = wb.create_sheet(title)
    banner(ws, banner_text, len(cols))
    header(ws, HDR, cols)
    for r, row in enumerate(rows, FIRST):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v).alignment = WRAP
    widths(ws, col_widths)
    ws.freeze_panes = ws.cell(FIRST, 1)
    return ws


def build(est, der, a):
    wb = Workbook()
    state = a.state
    bn = (f"DRAFT — failed checks: {a.draft}" if a.draft else f"Estimate state: {state}") + \
         f" | {a.client or est.get('lead_ref', '')} | version {a.version} | derived numbers: calculator"
    kinds = {o["kind"]: o for o in der["overheads"]}

    # ---------------- Estimation Breakdown
    ws = wb.active
    ws.title = EB
    banner(ws, bn, 13)
    header(ws, HDR, ["#", "Features", "Estimation Type", "Role", "Optimistic (hrs)", "Pessimistic (hrs)", "Most Likely (hrs)", "PERT (hrs)"])
    r, n = FIRST, 0
    for m in est["modules"]:
        band = r
        ws.cell(r, 2, m["module"]).font = BOLD
        for c in range(1, 9):
            ws.cell(r, c).fill = BFILL
        r += 1
        for t in m["tasks"]:
            n += 1
            ws.cell(r, 1, n)
            ws.cell(r, 2, f"{t['title']} — {t.get('delivers') or ''}".rstrip(" —")).alignment = WRAP
            ws.cell(r, 3, t["type"])
            ws.cell(r, 4, t["role"])
            ws.cell(r, 5, float(t["optimistic_h"]))
            ws.cell(r, 6, float(t["pessimistic_h"]))
            ws.cell(r, 7, float(t["most_likely_h"]))
            ws.cell(r, 8, f"=(E{r}+4*G{r}+F{r})/6")
            r += 1
        ws.cell(band, 8, f"=SUM(H{band + 1}:H{r - 1})").font = BOLD
    last = r - 1
    tot = r
    ws.cell(tot, 2, "Total").font = BOLD
    for col in "EFGH":
        ws[f"{col}{tot}"] = f'=SUMIF($A${FIRST}:$A${last},">0",{col}${FIRST}:{col}${last})'
        ws[f"{col}{tot}"].font = BOLD
    for rr in range(FIRST, tot + 1):
        for col in "EFGH":
            ws[f"{col}{rr}"].number_format = "0.0"

    # summary block (J:M)
    for i, h in enumerate(["Summary", "%", "PERT (hrs)", "Note"], 10):
        c = ws.cell(HDR, i, h)
        c.font, c.fill = BOLD, HFILL
    ws.cell(R_DEV, 10, "Development + Design")
    ws.cell(R_DEV, 12, f"=H{tot}")
    for k, row in ROW_OF.items():
        o = kinds.get(k)
        ws.cell(row, 10, LABELS[k])
        ws.cell(row, 11, o["pct"] if o else 0)
        ws.cell(row, 12, f"=$L${R_DEV}*K{row}/100")
        if o and o.get("reason_if_changed"):
            ws.cell(row, 13, o["reason_if_changed"]).alignment = WRAP
    ws.cell(R_TOT, 10, "Total").font = BOLD
    ws.cell(R_TOT, 12, f"=SUM(L{R_DEV}:L{R_CONT})").font = BOLD
    ws.cell(R_O, 10, "Optimistic total")
    ws.cell(R_O, 12, f"=E{tot}*(1+SUM(K{R_QA}:K{R_CONT})/100)")
    ws.cell(R_M, 10, "Most Likely total")
    ws.cell(R_M, 12, f"=G{tot}*(1+SUM(K{R_QA}:K{R_CONT})/100)")
    ws.cell(R_P, 10, "Pessimistic total")
    ws.cell(R_P, 12, f"=F{tot}*(1+SUM(K{R_QA}:K{R_CONT})/100)")
    ws.cell(R_P, 13, "Same overhead % applied to each column's base.")
    ws.cell(R_P80, 10, "P80 total (calculator)")
    ws.cell(R_P80, 12, der["total"]["p80"])
    ws.cell(R_P80, 13, "Static value: base + 0.8416 x sqrt(sum of task variances), scaled by overheads.")
    ws.cell(R_STATE, 10, "Estimate state")
    ws.cell(R_STATE, 12, state)
    ws.cell(R_HPW, 10, "Productive hours per FTE-week")
    ws.cell(R_HPW, 12, der["inputs"]["productive_hours_per_fte_week"])
    for rr in range(R_DEV, R_HPW + 1):
        ws.cell(rr, 12).number_format = "0.0"
    widths(ws, [5, 70, 24, 20, 12, 12, 12, 11, 3, 36, 7, 12, 50])
    ws.freeze_panes = ws.cell(FIRST, 1)

    # ---------------- Assumptions (UWS layout: three columns)
    asm = [f"{x['id']}: {x['text']}" for x in est.get("assumptions", [])] + \
          [f"Assumed gap: {g['gap']} — {g['assumption']}" for g in est.get("gaps_assumed", [])]
    ext = list(est.get("exclusions", []))
    dep = [f"{d['dependency']} (needed by {d['needed_by']}; if late: {d['impact_if_late']})" for d in est.get("client_dependencies", [])]
    dep += [f"Open question {q['id']}: {q.get('question', '')}" for q in est.get("open_questions", [])]
    rows = [(asm[i] if i < len(asm) else None, ext[i] if i < len(ext) else None, dep[i] if i < len(dep) else None)
            for i in range(max(len(asm), len(ext), len(dep), 1))]
    simple_sheet(wb, "Assumptions", bn, ["Assumptions", "Out of Scope", "Client Dependencies / Open Questions"], rows, [60, 45, 60])

    # ---------------- Estimation Basis
    cont = kinds.get("cont")
    ov_txt = lambda k, nm: f"{nm}: {kinds[k]['pct']}% of Dev + Design PERT" + (f" ({kinds[k]['reason_if_changed']})" if kinds[k].get("reason_if_changed") else "") if k in kinds else f"{nm}: not applied"
    conf = est.get("confidence", {})
    basis = [
        ("Estimation Method", "Three-point PERT per task: (Optimistic + 4 x Most Likely + Pessimistic) / 6. Overheads are percentages of the Dev + Design PERT base."),
        ("Source Documents", a.sources or "not recorded"),
        ("Functional Scope Basis", est.get("understanding", {}).get("scope_summary", "")),
        ("Task Size", "Most Likely target 30-40 h per task; larger tasks split unless atomic."),
        ("QA Basis", ov_txt("qa", "QA + security validation (excludes formal penetration testing)")),
        ("PM Basis", ov_txt("pm", "Project management")),
        ("UAT Basis", ov_txt("uat", "UAT support")),
        ("Documentation Basis", ov_txt("docs", "Documentation and knowledge transfer")),
        ("Contingency Basis", ov_txt("cont", "Contingency")),
        ("Productive hours", f"{der['inputs']['productive_hours_per_fte_week']} productive hours per FTE-week (assumed default)."),
        ("Timeline (calculator)", f"{der['timeline']['weeks_min']} / {der['timeline']['weeks_likely']} / {der['timeline']['weeks_max']} weeks (min / likely / max), contingency applied as schedule buffer."),
        ("Confidence", f"{conf.get('overall', '')} — {conf.get('reason', '')}"),
        ("Re-estimation Triggers", "; ".join(est.get("re_estimation_triggers", []))),
        ("Estimate state", state),
        ("Calculator checks", f"{sum(1 for c in der['checks'] if c['pass'])}/{len(der['checks'])} pass"),
    ]
    if der.get("ensemble"):
        basis.append(("Ensemble", f"Median of several runs; total spread {der['ensemble']['spread_pct']}% ({der['ensemble']['level']})."))
    simple_sheet(wb, "Estimation Basis", bn, ["Item", "Basis"], basis, [28, 120])

    # ---------------- Effort by Type
    ws = wb.create_sheet("Effort by Type")
    banner(ws, bn, 6)
    header(ws, HDR, ["Type", "Applicability", "Reason / note", "Task hrs", "Overhead hrs", "Total hrs"])
    appl = {x["type"]: x for x in est.get("type_applicability", [])}
    alloc = {"Testing / QA": [R_QA, R_UAT], "Project Management": [R_PM], "Documentation": [R_DOC]}
    r = FIRST
    for ty in ec.TYPES:
        ap = appl.get(ty, {})
        ws.cell(r, 1, ty)
        ws.cell(r, 2, ap.get("applicability"))
        note = ap.get("reason") or ""
        if ty == "Maintenance":
            note = (note + " " if note else "") + "Reported on the Maintenance sheet; never in the build total."
        ws.cell(r, 3, note).alignment = WRAP
        ws.cell(r, 4, f"=SUMIF('{EB}'!$C${FIRST}:$C${last},A{r},'{EB}'!$H${FIRST}:$H${last})")
        refs = alloc.get(ty, [])
        ws.cell(r, 5, "=" + "+".join(f"'{EB}'!$L${x}" for x in refs) if refs else 0)
        ws.cell(r, 6, f"=D{r}+E{r}")
        r += 1
    ws.cell(r, 1, "Contingency (separate line)")
    ws.cell(r, 6, f"='{EB}'!$L${R_CONT}")
    r += 1
    ws.cell(r, 1, "Total build").font = BOLD
    ws.cell(r, 6, f"=SUM(F{FIRST}:F{r - 1})").font = BOLD
    r += 1
    ws.cell(r, 1, "Difference vs Estimation Breakdown total (should be 0)")
    ws.cell(r, 6, f"=F{r - 1}-'{EB}'!$L${R_TOT}")
    for rr in range(FIRST, r + 1):
        for c in (4, 5, 6):
            ws.cell(rr, c).number_format = "0.0"
    widths(ws, [34, 16, 60, 12, 14, 12])
    ws.freeze_panes = ws.cell(FIRST, 1)

    # ---------------- Team & Timeline (static from calculator)
    ws = wb.create_sheet("Team & Timeline")
    banner(ws, bn, 6)
    header(ws, HDR, ["Role", "Hours", "FTE", "Phases", "", ""])
    r = FIRST
    for t in der["team"]:
        ws.cell(r, 1, t["role"])
        ws.cell(r, 2, t["hours"]).number_format = "0.0"
        ws.cell(r, 3, t["fte"])
        ws.cell(r, 4, ", ".join(t["phases"]) if t.get("phases") else "")
        r += 1
    r += 1
    ws.cell(r, 1, "Timeline (weeks): min / likely / max").font = BOLD
    ws.cell(r, 2, der["timeline"]["weeks_min"])
    ws.cell(r, 3, der["timeline"]["weeks_likely"])
    ws.cell(r, 4, der["timeline"]["weeks_max"])
    r += 2
    header(ws, r, ["Phase", "Start (wk)", "Weeks", "End (wk)", "Hours", "Capacity (h/wk)"])
    for name, p in der["phases"].items():
        r += 1
        for c, v in enumerate([name, p["start"], p["weeks"], p["end"], p["hours"], p["capacity_h_per_week"]], 1):
            ws.cell(r, c, v)
    widths(ws, [34, 12, 10, 40, 12, 16])

    # ---------------- remaining simple sheets
    simple_sheet(wb, "Tech Stack", bn, ["Layer", "Proposed", "Alternative", "Classification", "Rationale", "Effort impact"],
                 [(t["layer"], t["proposed"], t.get("alternative"), t.get("classification"), t.get("rationale"), t.get("effort_impact"))
                  for t in est.get("tech_stack", [])], [22, 30, 30, 18, 50, 40])
    simple_sheet(wb, "Risks", "INTERNAL — " + bn, ["ID", "Risk", "L", "I", "Effort exposure", "Mitigation", "Linked to"],
                 [(x["id"], x["risk"], x["likelihood"], x["impact"], x.get("effort_exposure"), x.get("mitigation"), ", ".join(x.get("linked_to", [])))
                  for x in est.get("risks", [])], [8, 60, 5, 5, 16, 50, 24])
    simple_sheet(wb, "Cost Drivers", bn, ["Item", "Category", "Driver", "One-time / recurring", "Sizing assumption", "Who bears cost", "Price status"],
                 [(c["item"], c["category"], c["driver"], c["one_time_or_recurring"], c.get("sizing_assumption"), c["who_bears_cost"], c["price_status"])
                  for c in est.get("cost_drivers", [])], [30, 22, 28, 18, 40, 14, 14])
    mr = est.get("maintenance_recommendation") or {}
    simple_sheet(wb, "Maintenance", bn, ["Item", "Value"],
                 [("Warranty / post-launch window (days)", mr.get("warranty_days")), ("AMC support hours per year", mr.get("amc_hours_per_year")),
                  ("Scope", "; ".join(mr.get("scope", []))), ("Reasoning", mr.get("reasoning")),
                  ("Note", "Hours only. Not part of the build total.")], [40, 100])
    clar = json.load(open(a.clarifications, encoding="utf-8")) if a.clarifications else []
    if not clar:
        clar = [{"id": q["id"], "question": "", "asked_by": "estimation-agent", "answer": q.get("answer"), "effect": q.get("effect_on_estimate")}
                for q in est.get("answered_questions", [])]
    simple_sheet(wb, "Clarifications Log", bn, ["Question id", "Question", "Asked by", "Answer", "Effect on estimate"],
                 [(q.get("id"), q.get("question"), q.get("asked_by"), q.get("answer"), q.get("effect")) for q in clar], [12, 60, 18, 50, 50])
    rc = [(i["id"], i["kind"], i["ref"], i["text"], i["coverage"], ", ".join(i.get("task_ids", [])), i.get("note")) for i in est.get("requirement_items", [])]
    rc += [("(inferred task)", "task", t["id"], f"{t['title']}", "inferred_necessary", t["id"], t.get("notes"))
           for m in est["modules"] for t in m["tasks"] if t.get("origin") == "inferred_necessary"]
    simple_sheet(wb, "Requirement Coverage", bn, ["Item", "Kind", "Source ref", "Text", "Coverage", "Tasks", "Note"], rc, [16, 14, 36, 60, 18, 18, 40])
    return wb, {"tasks": n, "first": FIRST, "last": last, "total_row": tot}


def check(path, der):
    """Independent recomputation from the workbook's input cells (does not evaluate formulas)."""
    wb = load_workbook(path)
    ws = wb[EB]
    O = P = M = pert = 0.0
    n = 0
    for row in ws.iter_rows(min_row=FIRST, max_col=7):
        if isinstance(row[0].value, (int, float)):
            o, p, m = (float(row[i].value) for i in (4, 5, 6))
            O, P, M, pert, n = O + o, P + p, M + m, pert + (o + 4 * m + p) / 6, n + 1
    pct = sum(float(ws.cell(rr, 11).value or 0) for rr in (R_QA, R_PM, R_UAT, R_DOC, R_CONT))
    tot = pert * (1 + pct / 100)
    ok = n == der["inputs"]["task_count"] and abs(tot - der["total"]["pert"]) <= 0.15 and abs(pert - der["base"]["pert"]) <= 0.15
    print(f"check: tasks {n} (derived {der['inputs']['task_count']}), base {pert:.1f} (derived {der['base']['pert']}), "
          f"total {tot:.1f} (derived {der['total']['pert']}) -> {'OK' if ok else 'MISMATCH'}")
    return ok


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("estimate")
    ap.add_argument("derived")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--client", default="")
    ap.add_argument("--version", default="1.0")
    ap.add_argument("--state", default="draft", choices=["draft", "delivery_reviewed", "approved"])
    ap.add_argument("--draft", default="", help="comma-separated failed check ids; stamps the workbook DRAFT")
    ap.add_argument("--clarifications")
    ap.add_argument("--sources", default="")
    ap.add_argument("--check", action="store_true")
    a = ap.parse_args()
    est = ec.load(a.estimate)
    der = json.load(open(a.derived, encoding="utf-8"))
    wb, info = build(est, der, a)
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    print(f"wrote {a.out}: {info['tasks']} tasks, {len(wb.sheetnames)} sheets")
    if a.check and not check(a.out, der):
        sys.exit(1)


if __name__ == "__main__":
    main()
